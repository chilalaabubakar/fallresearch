"""Build the Applied Energy manuscript (Word) from the Markdown sources.

    python docs/paper/build.py

* Citations [@key] or [@a; @b] become numbers in order of first appearance,
  e.g. [1] or [2,3] or [4–6]; the list replaces {{REFERENCES}} (references.yaml).
* Caption blocks "> **Fig. N.** ..." become the image with its caption below.
* {{TABLE_S1}} in the supplementary material becomes the parameter table.
* Checks: 3-5 highlights of at most 85 characters, abstract of at most 250
  words; prints the word count and the references still marked "check".
Outputs: paper.docx, supplementary.docx, highlights.docx and cover_letter.docx in docs/paper/.
"""

import re
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
import yaml
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIGURE_WIDTH = "16cm"
# Figure label -> image (results/paper/, made by python -m topside.figures)
FIGURES = {"1": "fig00_system", "2": "fig02_demand_map", "3": "fig04_cases", "4": "fig05_example_week",
           "5": "fig06_gt_load_duration", "6": "fig07_injection_mechanism",
           "7": "fig09_sensitivity_overview", "8": "fig10_validation", "9": "fig11_multi_year",
           "S1": "fig01_wind_resource", "S2": "fig03_gt_part_load", "S3": "fig08_monthly_co2"}
MAX_HIGHLIGHT_CHARS, HIGHLIGHTS = 85, (3, 5)
MAX_ABSTRACT_WORDS = 250
CITE = re.compile(r"\[(@[\w-]+(?:\s*;\s*@[\w-]+)*)\]")


def compress(numbers):
    """[1, 2, 3, 5] -> '1–3,5'."""
    numbers, out, i = sorted(set(numbers)), [], 0
    while i < len(numbers):
        j = i
        while j + 1 < len(numbers) and numbers[j + 1] == numbers[j] + 1:
            j += 1
        out.append(str(numbers[i]) if j - i < 2 else f"{numbers[i]}–{numbers[j]}")
        if j - i == 1:
            out.append(str(numbers[j]))
        i = j + 1
    return ",".join(out)


def number_citations(markdown, refs):
    order = []

    def repl(m):
        keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        for k in keys:
            if k not in refs:
                raise KeyError(f"unknown reference [@{k}]")
            if k not in order:
                order.append(k)
        return "\\[" + compress([order.index(k) + 1 for k in keys]) + "\\]"

    markdown = CITE.sub(repl, markdown)
    listing = "\n\n".join(f"\\[{i}\\] {refs[k]['text']}" for i, k in enumerate(order, 1))
    unused = sorted(set(refs) - set(order))
    return markdown.replace("{{REFERENCES}}", listing), order, unused


def insert_figures(markdown):
    def repl(m):
        block = re.sub(r"^>\s?", "", m.group(0), flags=re.M).strip()
        label = re.match(r"\*\*Fig\. (S?\d+)\.\*\*", block).group(1)
        path = (ROOT / "results" / "paper" / f"{FIGURES[label]}.png").relative_to(ROOT).as_posix()
        if not (ROOT / path).exists():
            raise FileNotFoundError(f"{path} (run python -m topside.figures)")
        return f"![{' '.join(block.split())}]({path}){{width={FIGURE_WIDTH}}}\n"
    return re.sub(r"(?m)^> \*\*Fig\. S?\d+\.\*\*.*(?:\n>.*)*", repl, markdown)


def parameter_table():
    t = pd.read_csv(ROOT / "results" / "paper" / "tables" / "T1_parameters.csv")
    t = t[["parameter", "value", "unit", "source"]].fillna("")
    for c in t:
        t[c] = t[c].astype(str).str.replace("|", "/", regex=False)
    return t.to_markdown(index=False)


def add_table_borders(docx_path):
    doc = Document(docx_path)
    for table in doc.tables:
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "808080")):
                el.set(qn(f"w:{k}"), v)
            borders.append(el)
        table._tbl.tblPr.append(borders)
    doc.save(docx_path)


def to_docx(markdown, out):
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as tmp:
        tmp.write(markdown)
    subprocess.run(["pandoc", tmp.name, "-f", "markdown", "-t", "docx", "-o", str(out),
                    "--resource-path", str(ROOT)], check=True)
    Path(tmp.name).unlink()
    add_table_borders(out)
    return out


def words(text):
    text = re.sub(r"(?m)^\|.*$", "", text)            # tables
    text = re.sub(r"(?m)^>.*$", "", text)             # figure captions
    text = re.sub(r"\$\$.*?\$\$", "", text, flags=re.S)
    return len(re.findall(r"[A-Za-z0-9][\w'’–-]*", text))


def section(md, start, end):
    return md[md.index(start): md.index(end)]


def check_highlights(md):
    items = [l[2:].strip() for l in md.splitlines() if l.startswith("* ")]
    problems = [f"{len(h)} chars: {h}" for h in items if len(h) > MAX_HIGHLIGHT_CHARS]
    if not HIGHLIGHTS[0] <= len(items) <= HIGHLIGHTS[1]:
        problems.append(f"{len(items)} highlights (need {HIGHLIGHTS[0]}-{HIGHLIGHTS[1]})")
    return items, problems


def main():
    refs = yaml.safe_load((HERE / "references.yaml").read_text(encoding="utf-8"))
    paper = (HERE / "paper.md").read_text(encoding="utf-8")
    n_abstract = words(section(paper, "## Abstract", "**Keywords:**")) - 1
    n_body = words(section(paper, "## 1 Introduction", "## Declarations"))
    paper, order, unused = number_citations(paper, refs)
    print("Saved", to_docx(insert_figures(paper), HERE / "paper.docx"))

    supp = (HERE / "supplementary.md").read_text(encoding="utf-8").replace("{{TABLE_S1}}", parameter_table())
    print("Saved", to_docx(insert_figures(supp), HERE / "supplementary.docx"))

    letter = HERE / "cover_letter.md"
    if letter.exists():
        print("Saved", to_docx(letter.read_text(encoding="utf-8"), HERE / "cover_letter.docx"))

    hl = (HERE / "highlights.md").read_text(encoding="utf-8")
    items, problems = check_highlights(hl)
    print("Saved", to_docx(hl, HERE / "highlights.docx"))

    print(f"\nAbstract: {n_abstract} words (limit {MAX_ABSTRACT_WORDS})"
          + ("  ** TOO LONG **" if n_abstract > MAX_ABSTRACT_WORDS else ""))
    print(f"Main text (Introduction to Conclusions, without tables and captions): {n_body} words")
    print(f"Highlights: {len(items)}, longest {max(map(len, items))} characters"
          + "".join(f"\n  ** {p}" for p in problems))
    print(f"References cited: {len(order)}" + (f"; not cited: {', '.join(unused)}" if unused else ""))
    to_check = [f"[{order.index(k) + 1}] {k}" for k in order if refs[k].get("check")]
    if to_check:
        print("References to confirm against the original:\n  " + "\n  ".join(to_check))


if __name__ == "__main__":
    main()
