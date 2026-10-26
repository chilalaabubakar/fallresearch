"""Build Word versions of the thesis chapters from the Markdown drafts.

    python docs/thesis/build_word.py

* Each figure caption block ("> **Figure N.** ...") is replaced by the image
  results/paper/figNN_*.png with the caption below it.
* Equations ($$ ... $$) become native Word equations (pandoc OMML).
* Tables get single-line borders (pandoc leaves them borderless).
* comparison.md is appended to the results chapter as Section 9 (headings
  demoted one level, its italic draft note dropped).
Outputs: docs/thesis/methods.docx and docs/thesis/results.docx
"""

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIGURES = ROOT / "results" / "paper"
CHAPTERS = ["methods", "results"]
# Markdown files appended to a chapter as numbered sections: chapter -> [(file, section number)]
SECTIONS = {"results": [("comparison", 9)]}
# Figures that are not results/paper/figNN_*.png (site comparison chapter)
FIGURE_FILES = {11: "results/comparison/site_comparison.png",
                12: "results/tanzania/paper/fig05_example_week.png",
                13: "results/tanzania/paper/fig09_sensitivity_overview.png"}
FIGURE_WIDTH = "16cm"


def figure_path(number):
    if number in FIGURE_FILES:
        return FIGURE_FILES[number]
    matches = sorted(FIGURES.glob(f"fig{number:02d}_*.png"))
    if not matches:
        raise FileNotFoundError(f"No image for figure {number} in {FIGURES}")
    return matches[0].relative_to(ROOT).as_posix()


def insert_figures(markdown):
    """Replace '> **Figure N.** caption' blockquotes with an image + caption."""
    def repl(m):
        block = re.sub(r"^>\s?", "", m.group(0), flags=re.M).strip()
        number = int(re.match(r"\*\*Figure (\d+)\.\*\*", block).group(1))
        caption = " ".join(block.split())
        return f"![{caption}]({figure_path(number)}){{width={FIGURE_WIDTH}}}\n"
    return re.sub(r"(?m)^> \*\*Figure \d+\.\*\*.*(?:\n>.*)*", repl, markdown)


def add_table_borders(docx_path):
    doc = Document(docx_path)
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        tbl_pr.append(borders)
    doc.save(docx_path)


def as_section(markdown, number):
    """Turn a stand-alone chapter into section `number` of another chapter."""
    markdown = re.sub(r"\A# .*\n+\*[^*]+\*\n+", lambda m: m.group(0).split("\n")[0] + "\n\n",
                      markdown, count=1)                       # drop the italic draft note
    markdown = re.sub(r"(?m)^## (\d+) ", lambda m: f"### {number}.{m.group(1)} ", markdown)
    return re.sub(r"(?m)^# ", f"## {number} ", markdown, count=1)


def build(chapter):
    src = HERE / f"{chapter}.md"
    out = HERE / f"{chapter}.docx"
    md = src.read_text(encoding="utf-8")
    for name, number in SECTIONS.get(chapter, []):
        md = md.rstrip() + "\n\n" + as_section((HERE / f"{name}.md").read_text(encoding="utf-8"), number)
    md = insert_figures(md)
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as tmp:
        tmp.write(md)
    subprocess.run(["pandoc", tmp.name, "-f", "markdown", "-t", "docx", "-o", str(out),
                    "--resource-path", str(ROOT)], check=True)
    Path(tmp.name).unlink()
    add_table_borders(out)
    return out


if __name__ == "__main__":
    for c in CHAPTERS:
        print("Saved", build(c))
